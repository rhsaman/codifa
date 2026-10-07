"""Document tool round-trips: create → read → edit for docx / xlsx / pdf / text.

Locks in the read_document / write_document contract the model relies on:
markdown-lite rendering, find/replace, cell sets, row appends and PDF page
ops — all offline, all hermetic (files land in tmp_path).
"""

import json
import os
import sys

_THIS = os.path.dirname(os.path.abspath(__file__))
for _p in (_THIS, os.path.dirname(_THIS)):
    if _p not in sys.path:
        sys.path.insert(0, _p)

import pytest

import doc_tools
from doc_tools import (
    DocError,
    create_document,
    edit_document,
    read_document,
    style_document,
)

# -- text -------------------------------------------------------------------

def test_text_create_and_read_roundtrip(tmp_path):
    p = str(tmp_path / "notes.md")
    out = create_document(p, "# Title\n\nbody line")
    assert os.path.exists(p)
    assert "created" in out
    assert read_document(p) == "# Title\n\nbody line"


def test_create_requires_content(tmp_path):
    with pytest.raises(DocError):
        create_document(str(tmp_path / "a.txt"), "  ")


def test_create_makes_parent_dirs(tmp_path):
    p = str(tmp_path / "deep" / "nest" / "a.md")
    create_document(p, "hi")
    assert os.path.exists(p)


# -- Word -------------------------------------------------------------------

def test_docx_create_read_roundtrip(tmp_path):
    p = str(tmp_path / "report.docx")
    create_document(
        p,
        "# Report\n\nfirst paragraph with **bold** text\n\n- bullet one\n- bullet two",
    )
    text = read_document(p)
    assert "# Report" in text
    assert "first paragraph with bold text" in text  # bold marker stripped on read
    assert "- bullet one" in text
    assert "- bullet two" in text


def test_docx_persian_roundtrip(tmp_path):
    p = str(tmp_path / "fa.docx")
    create_document(p, "# گزارش\n\nمتن آزمایشی فارسی")
    text = read_document(p)
    assert "گزارش" in text
    assert "متن آزمایشی فارسی" in text


def test_docx_find_replace(tmp_path):
    p = str(tmp_path / "r.docx")
    create_document(p, "the quick brown fox")
    out = edit_document(p, find="brown", replace="red")
    assert "replaced 1 occurrence" in out
    assert "the quick red fox" in read_document(p)
    assert "brown" not in read_document(p)


def test_docx_replace_across_runs(tmp_path):
    # "brown" split across two runs must still be replaced (paragraph-level
    # fallback path).
    from docx import Document

    p = str(tmp_path / "runs.docx")
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("the quick bro")
    para.add_run("wn fox")
    doc.save(p)
    edit_document(p, find="brown", replace="red")
    assert "the quick red fox" in read_document(p)


def test_docx_append_content(tmp_path):
    p = str(tmp_path / "app.docx")
    create_document(p, "# Head\n\nbody")
    out = edit_document(p, content="# Extra\n\nmore text")
    assert "appended" in out
    text = read_document(p)
    assert "# Extra" in text
    assert "more text" in text


def test_docx_edit_needs_args(tmp_path):
    p = str(tmp_path / "x.docx")
    create_document(p, "hello")
    with pytest.raises(DocError):
        edit_document(p)


def test_docx_edit_zero_match_is_an_error(tmp_path):
    """A find/replace that matches nothing must FAIL — the model must not be
    able to tell the user 'done' while nothing changed."""
    p = str(tmp_path / "z.docx")
    create_document(p, "the quick fox")
    with pytest.raises(DocError, match="not found"):
        edit_document(p, find="brown", replace="red")
    # And the file is untouched.
    assert "the quick fox" in read_document(p)


# -- rich rendering for the viewer -------------------------------------------

def test_docx_render_html(tmp_path):
    p = str(tmp_path / "r.docx")
    create_document(p, "# قرارداد\n\nمتن با **بولد**\n\n- گزینه یک\n- گزینه دو")
    html = doc_tools.render_document(p)
    assert '<h1 dir="auto">قرارداد</h1>' in html
    assert "<li" in html and "گزینه یک" in html
    assert "font-weight:700" in html
    assert "<table" not in html


def test_docx_render_html_carries_fonts(tmp_path):
    p = str(tmp_path / "f.docx")
    create_document(p, "plain text")
    style_document(p, font="Vazirmatn", size=18)
    html = doc_tools.render_document(p)
    assert "font-family:Vazirmatn" in html
    assert "font-size:18pt" in html


def test_xlsx_render_html(tmp_path):
    p = str(tmp_path / "x.xlsx")
    create_document(p, json.dumps([["نام", "مبلغ"], ["الف", 12]]))
    style_document(p, bold=True, rng="A1:B1")
    html = doc_tools.render_document(p)
    assert "<h3" in html and "<table" in html
    assert "نام" in html and "12" in html
    assert "font-weight:700" in html


def test_render_pdf_pages(tmp_path):
    p = str(tmp_path / "p.pdf")
    create_document(p, "hello page\n\n---\n\nsecond page")
    pages = doc_tools.render_pdf_pages(p)
    assert len(pages) == 2
    assert all(u.startswith("data:image/jpeg;base64,") for u in pages)


def test_render_pdf_pages_capped(tmp_path):
    p = str(tmp_path / "many.pdf")
    create_document(p, "\n\n---\n\n".join(f"page {i}" for i in range(20)))
    pages = doc_tools.render_pdf_pages(p)
    assert len(pages) == doc_tools.MAX_PDF_PREVIEW_PAGES


# -- styling ------------------------------------------------------------------

def test_docx_style_whole_document(tmp_path):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import RGBColor

    p = str(tmp_path / "s.docx")
    create_document(p, "# Head\n\nplain text\n\n- bullet")
    out = style_document(p, font="Vazirmatn", size=14, bold=True, color="FF0000")
    assert "styled" in out and "whole document" in out
    doc = Document(p)
    runs = [r for para in doc.paragraphs for r in para.runs if r.text.strip()]
    assert runs
    assert all(r.font.name == "Vazirmatn" for r in runs)
    assert all(r.font.size.pt == 14 for r in runs)
    assert all(r.font.bold for r in runs)
    assert all(r.font.color.rgb == RGBColor(0xFF, 0x00, 0x00) for r in runs)
    # The complex-script slot must carry the family too — Word renders
    # Persian/Arabic from it, otherwise font changes do nothing on RTL text.
    assert all(
        r._element.rPr.rFonts.get(qn("w:cs")) == "Vazirmatn" for r in runs
    )


def test_docx_style_with_find_filter(tmp_path):
    from docx import Document

    p = str(tmp_path / "f.docx")
    create_document(p, "# Keep\n\nchange me please")
    out = style_document(p, font="Arial", find="change me")
    assert "matching 'change me'" in out
    doc = Document(p)
    by_text = {
        r.text: r.font.name
        for para in doc.paragraphs
        for r in para.runs
        if r.text.strip()
    }
    assert by_text.get("change me please") == "Arial"
    assert by_text.get("Keep") != "Arial"


def test_docx_style_no_match_errors(tmp_path):
    p = str(tmp_path / "n.docx")
    create_document(p, "hello")
    with pytest.raises(DocError, match="nothing styled"):
        style_document(p, font="Arial", find="zzz")


def test_docx_style_needs_args(tmp_path):
    p = str(tmp_path / "a.docx")
    create_document(p, "hello")
    with pytest.raises(DocError):
        style_document(p)


def test_xlsx_style_range(tmp_path):
    import openpyxl

    p = str(tmp_path / "st.xlsx")
    create_document(p, json.dumps([["a", "b"], ["c", "d"]]))
    out = style_document(p, font="Tahoma", size=12, bold=True, rng="A1:B2")
    assert "styled 4 cell(s)" in out
    wb = openpyxl.load_workbook(p)
    for row in wb.active["A1:B2"]:
        for c in row:
            assert c.font.name == "Tahoma"
            assert c.font.size == 12
            assert c.font.bold


def test_xlsx_style_single_cell_and_color(tmp_path):
    import openpyxl

    p = str(tmp_path / "sc.xlsx")
    create_document(p, json.dumps([["v"]]))
    style_document(p, color="FF0000", cell="A1")
    wb = openpyxl.load_workbook(p)
    assert wb.active["A1"].font.color.rgb == "00FF0000"


def test_xlsx_style_missing_sheet_errors(tmp_path):
    p = str(tmp_path / "ms.xlsx")
    create_document(p, json.dumps([["a"]]))
    with pytest.raises(DocError, match="not found"):
        style_document(p, font="Arial", sheet="Nope")


def test_style_pdf_rejected(tmp_path):
    p = str(tmp_path / "x.pdf")
    create_document(p, "hello")
    with pytest.raises(DocError, match="recreate"):
        style_document(p, font="Arial")


def test_doc_read_rejects_legacy_doc(tmp_path):
    p = str(tmp_path / "legacy.doc")
    with open(p, "w") as fh:
        fh.write("whatever")
    with pytest.raises(DocError):
        read_document(p)


# -- Excel ------------------------------------------------------------------

def test_xlsx_create_read_roundtrip_json(tmp_path):
    p = str(tmp_path / "table.xlsx")
    rows = [["name", "qty"], ["widget", 3], ["gadget", 7]]
    create_document(p, json.dumps(rows))
    text = read_document(p)
    assert "### sheet: Sheet1" in text
    assert "name\tqty" in text
    assert "widget\t3" in text
    assert "gadget\t7" in text


def test_xlsx_create_from_json_objects(tmp_path):
    p = str(tmp_path / "obj.xlsx")
    create_document(p, json.dumps([{"a": 1, "b": "x"}, {"a": 2, "b": "y"}]))
    text = read_document(p)
    assert "a\tb" in text
    assert "2\ty" in text


def test_xlsx_create_from_tsv(tmp_path):
    p = str(tmp_path / "tsv.xlsx")
    create_document(p, "name\tqty\nwidget\t3")
    text = read_document(p)
    assert "widget\t3" in text


def test_xlsx_set_cell_and_append_rows(tmp_path):
    p = str(tmp_path / "edit.xlsx")
    create_document(p, json.dumps([["h1", "h2"], ["a", "b"]]))
    out = edit_document(p, cell="B2", value="99")
    assert "set Sheet1!B2" in out
    text = read_document(p)
    assert "a\t99" in text
    out = edit_document(p, content="c\td", sheet="Sheet1")
    assert "appended 1 row(s)" in out
    assert "c\td" in read_document(p)


def test_xlsx_edit_creates_missing_sheet(tmp_path):
    p = str(tmp_path / "sh.xlsx")
    create_document(p, json.dumps([["a"]]))
    edit_document(p, content="x\ny", sheet="Extra")
    text = read_document(p, sheet="Extra")
    assert "### sheet: Extra" in text
    assert "x" in text and "y" in text


def test_xlsx_cell_value_coerced_to_number(tmp_path):
    import openpyxl

    p = str(tmp_path / "num.xlsx")
    create_document(p, json.dumps([["v"]]))
    edit_document(p, cell="A1", value="42")
    wb = openpyxl.load_workbook(p)
    assert wb.active["A1"].value == 42
    edit_document(p, cell="A1", value="4.5")
    wb = openpyxl.load_workbook(p)
    assert wb.active["A1"].value == 4.5


def test_xlsx_edit_needs_args(tmp_path):
    p = str(tmp_path / "e.xlsx")
    create_document(p, json.dumps([["a"]]))
    with pytest.raises(DocError):
        edit_document(p)


# -- PDF --------------------------------------------------------------------

def test_pdf_create_read_roundtrip_latin(tmp_path):
    p = str(tmp_path / "doc.pdf")
    out = create_document(p, "# Invoice\n\nTotal due: 42 EUR\n\n- item one\n- item two")
    assert "created PDF" in out
    text = read_document(p)
    assert "1 page" in text or "1 page(s)" in text
    assert "Invoice" in text
    assert "Total due: 42 EUR" in text


def test_pdf_create_persian_produces_pages(tmp_path):
    # Persian text must at least produce a parseable PDF with extractable
    # (reshaped) content — exact glyph round-trip depends on the font found.
    p = str(tmp_path / "fa.pdf")
    create_document(p, "# گزارش\n\nسلام دنیا")
    text = read_document(p)
    assert "page 1" in text
    assert "no extractable text" not in text.split("=== page 1 ===")[1].split("===")[0]


def test_pdf_extract_pages(tmp_path):
    p = str(tmp_path / "multi.pdf")
    create_document(
        p, "\n\n---\n\n".join(f"# Page {i}\ncontent {i}" for i in range(1, 6))
    )
    dest = str(tmp_path / "out.pdf")
    out = edit_document(p, pages="2-3", dest=dest)
    assert "extracted pages [2-3]" in out
    assert os.path.exists(dest)
    text = read_document(dest)
    assert "content 2" in text and "content 3" in text
    assert "content 4" not in text


def test_pdf_merge(tmp_path):
    a = str(tmp_path / "a.pdf")
    b = str(tmp_path / "b.pdf")
    create_document(a, "alpha document")
    create_document(b, "beta document")
    out = edit_document(a, source=b)
    assert "merged" in out
    merged = read_document(a)
    assert "alpha document" in merged
    assert "beta document" in merged


def test_pdf_delete_pages(tmp_path):
    p = str(tmp_path / "del.pdf")
    create_document(
        p, "\n\n---\n\n".join(f"# P{i}\ncontent {i}" for i in range(1, 4))
    )
    out = edit_document(p, pages="2")
    assert "deleted pages [2]" in out
    text = read_document(p)
    assert "content 1" in text
    assert "content 2" not in text
    assert "content 3" in text


def test_pdf_delete_all_pages_refused(tmp_path):
    p = str(tmp_path / "one.pdf")
    create_document(p, "only page")
    with pytest.raises(DocError):
        edit_document(p, pages="1")


def test_pdf_invalid_page_spec(tmp_path):
    p = str(tmp_path / "spec.pdf")
    create_document(p, "hello")
    with pytest.raises(DocError):
        edit_document(p, pages="abc", dest=str(tmp_path / "x.pdf"))


# -- plain read tool auto-extraction ----------------------------------------

def test_read_file_auto_extracts_documents(tmp_path):
    """The plain read tool must transparently extract PDF/docx text — the
    model should never be blocked by a 'binary file (read skipped)' error
    just because it picked `read` instead of `read_document`."""
    from tools import read_file

    root = str(tmp_path)
    create_document(os.path.join(root, "doc.pdf"), "hello auto extraction")
    res = read_file(root, "doc.pdf")
    assert "error" not in res
    assert "hello auto extraction" in res["content"]
    assert res.get("note") == "text auto-extracted from the document"

    create_document(os.path.join(root, "doc.docx"), "# Head\n\nbody text")
    res = read_file(root, "doc.docx")
    assert "error" not in res
    assert "body text" in res["content"]


# -- table parsing ----------------------------------------------------------

def test_parse_rows_rejects_bad_json(tmp_path):
    with pytest.raises(DocError):
        doc_tools._parse_rows("[not json")


def test_parse_rows_rejects_empty(tmp_path):
    with pytest.raises(DocError):
        doc_tools._parse_rows("   ")


# -- unknown / text-extension edits route to edit_file ----------------------

def test_edit_text_ext_routes_to_edit_file_hint(tmp_path):
    p = str(tmp_path / "note.md")
    create_document(p, "hello")
    with pytest.raises(DocError, match="edit_file"):
        edit_document(p, find="h", replace="H")
